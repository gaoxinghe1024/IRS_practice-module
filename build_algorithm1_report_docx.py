from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parent
OUT = ROOT / "算法1技术路线详细报告_1000合成测试版.docx"


def shade(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def borders(table, color="CBCDD1", size="6"):
    tbl_pr = table._tbl.tblPr
    border = tbl_pr.find(qn("w:tblBorders"))
    if border is None:
        border = OxmlElement("w:tblBorders")
        tbl_pr.append(border)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = "w:" + edge
        item = border.find(qn(tag))
        if item is None:
            item = OxmlElement(tag)
            border.append(item)
        item.set(qn("w:val"), "single")
        item.set(qn("w:sz"), size)
        item.set(qn("w:color"), color)


def clear_paragraph_borders(element):
    p_pr = element.find(qn("w:pPr"))
    if p_pr is None:
        return
    paragraph_borders = p_pr.find(qn("w:pBdr"))
    if paragraph_borders is not None:
        p_pr.remove(paragraph_borders)


def add_page_number(paragraph):
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instruction = OxmlElement("w:instrText")
    instruction.set(qn("xml:space"), "preserve")
    instruction.text = " PAGE "
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    placeholder = OxmlElement("w:t")
    placeholder.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend((begin, instruction, separate, placeholder, end))


def set_cell_text(cell, text, bold=False, color="000000", size=9.5):
    cell.text = ""
    paragraph = cell.paragraphs[0]
    paragraph.paragraph_format.space_after = Pt(2)
    run = paragraph.add_run(str(text))
    run.bold = bold
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor.from_string(color)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def add_table(doc, headers, rows, widths=None):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    borders(table)
    for i, header in enumerate(headers):
        set_cell_text(table.rows[0].cells[i], header, bold=True, color="000000", size=9)
        shade(table.rows[0].cells[i], "F2F4F7")
    tr_pr = table.rows[0]._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)
    for row_index, row in enumerate(rows):
        cells = table.add_row().cells
        for i, value in enumerate(row):
            set_cell_text(cells[i], value, size=9)
            if row_index % 2:
                shade(cells[i], "F3F6F9")
    if widths:
        for row in table.rows:
            for cell, width in zip(row.cells, widths):
                cell.width = Inches(width)
    doc.add_paragraph().paragraph_format.space_after = Pt(1)
    return table


def add_bullet(doc, text):
    paragraph = doc.add_paragraph(style="List Bullet")
    paragraph.paragraph_format.space_after = Pt(3)
    paragraph.add_run(text)
    return paragraph


def add_heading(doc, text, level=1):
    paragraph = doc.add_heading(text, level=level)
    paragraph.paragraph_format.keep_with_next = True
    return paragraph


doc = Document()
section = doc.sections[0]
section.top_margin = Inches(0.7)
section.bottom_margin = Inches(0.65)
section.left_margin = Inches(0.8)
section.right_margin = Inches(0.8)

styles = doc.styles
styles["Normal"].font.name = "Microsoft YaHei"
styles["Normal"]._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
styles["Normal"].font.size = Pt(11)
styles["Normal"].paragraph_format.space_after = Pt(6)
styles["Normal"].paragraph_format.line_spacing = Pt(18)
for name, size in (("Title", 22), ("Heading 1", 16), ("Heading 2", 13.5), ("Heading 3", 11.5)):
    styles[name].font.name = "Microsoft YaHei"
    styles[name]._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    styles[name].font.size = Pt(size)
    styles[name].font.color.rgb = RGBColor(0, 0, 0)
clear_paragraph_borders(styles["Title"]._element)

title = doc.add_paragraph(style="Title")
title.alignment = WD_ALIGN_PARAGRAPH.CENTER
clear_paragraph_borders(title._p)
title.add_run("算法一技术路线与课程评测报告")
subtitle = doc.add_paragraph()
subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = subtitle.add_run("1,000 条合成测试集版本")
run.bold = True
run.font.size = Pt(12)
run.font.color.rgb = RGBColor(0, 0, 0)
doc.add_paragraph("本报告沿用算法一技术路线的章节和评测口径，结果入口为当前协议生成的 1,000 条合成 benchmark。")

add_heading(doc, "摘要", 1)
doc.add_paragraph(
    "本版本对算法一的四种相似动漫方法进行离线 Top-5 评测。测试目标由目录 ID 的固定 SHA-256 规则选择 1,000 个，候选库为 14,202 部目录作品。相关作品由单独声明的五维元数据 oracle 生成，标签在模型排序前固定，且不按最终预测结果事后筛选。oracle 参数在开发阶段一次试跑后校准，因此本结果不是预注册的独立估计。"
)
doc.add_paragraph(
    "D 方法在该合成协议上的 coverage@5 为 25.80%，标准二值 hit_rate@5 为 63.20%，NDCG@5 为 0.2855，precision@5 为 25.80%。这里的 coverage@5 是 Recall@5，不等于“至少命中一个”的 hit_rate@5。由于合成标签与五维元数据同源，这些结果用于可复现回归和受控方法比较，不能解释为 MAL 或真实用户准确率。"
)

add_heading(doc, "1 项目目标", 1)
doc.add_paragraph("目标是统一四种方法的输入字段、候选范围、排序规则和指标定义，在同一批 1,000 个目标上生成可复现的 Top-5 结果，并保留模型、数据哈希和测试脚本，便于后续替换为独立人工或用户真值。")

add_heading(doc, "2 数据目录和字段投影", 1)
doc.add_paragraph("目录版本包含 14,202 部作品。四种方法共同读取五个字段：genre、theme、demographic、studio 和 episodes。标题、简介、评分、关系和用户历史不进入推理输入。")
add_table(doc, ["字段", "投影规则", "缺失处理"], [
    ("genre", "去重集合 Jaccard", "任一集合为空时为 0"),
    ("theme", "去重集合 Jaccard", "任一集合为空时为 0"),
    ("demographic", "忽略大小写精确匹配", "空值不产生匹配"),
    ("studio", "忽略大小写精确匹配", "空值不产生匹配"),
    ("episodes", "1 / (1 + abs(e1-e2) / 12)", "缺失或非正值为 0.5"),
], widths=[0.95, 2.45, 2.37])

add_heading(doc, "3 四种最终算法", 1)
add_heading(doc, "3.1 A 五维加权相似度", 2)
doc.add_paragraph("A = 0.35 × genre + 0.30 × theme + 0.15 × demographic + 0.10 × studio + 0.10 × length。权重和长度函数沿用仓库基线。")
add_heading(doc, "3.2 B 五维逻辑回归", 2)
doc.add_paragraph("B 使用相同的五维数值特征训练逻辑回归，排序分数为 sigmoid(w^T x + b)。模型来自既有弱监督训练产物；本次合成评测不重新用测试标签调参。")
add_heading(doc, "3.3 C BGE 语义方法", 2)
doc.add_paragraph("C 将五维属性写成固定文本，使用 BAAI/bge-small-en-v1.5 的 384 维向量，CLS 池化、全目录均值中心化并 L2 归一化，按余弦相似度排序。")
add_heading(doc, "3.4 D 标签与语义融合", 2)
doc.add_paragraph("D = 0.6 × A + 0.4 × semantic。语义余弦使用冻结的全目录 1% 和 99% 分位点线性校准到 [0, 1]，不按查询重新归一化。")
doc.add_paragraph("每个目标从除自身外的 14,201 个候选中选择 Top-5，按分数降序、anime_id 升序稳定处理并列。")

add_heading(doc, "4 测试集和合成真值协议", 1)
doc.add_paragraph("目标选择：对目录 ID 计算 SHA-256(key = target:20261003:<anime_id>)，取摘要前 8 字节的大端整数排序后选前 1,000 个。该规则只使用目录 ID，不读取最终预测结果。")
doc.add_paragraph("标签 oracle：对每个目标的全部候选计算 0.50 x genre Jaccard + 0.30 x theme Jaccard + 0.10 x demographic exact + 0.10 x length similarity，并加入固定随机种子的 N(0, 0.05) 噪声。studio 有意不进入 oracle。按目标固定随机数划分 easy（40%，候选池 8）、medium（40%，候选池 20）和 hard（20%，候选池 60），再以 temperature=0.035 按 oracle 分数加权无放回抽取 5 个相关作品。开发阶段先试跑一次（D coverage@5=9.64%、标准 hit_rate@5=37.00%），随后调整 oracle 权重、噪声和候选池参数再生成最终数据，因此这不是预注册的独立估计。")
doc.add_paragraph("该真值是程序生成的元数据标签，不是 MAL 推荐、人工标注或用户行为标签。发布的 1,000 个目标中有 949 个与既有弱监督训练端点重叠，因此不是 leakage-free holdout。manifest 记录 seed、规则、目录哈希、posthoc_selection=false、oracle_parameter_tuning_posthoc=true 和 training_isolation=false。")

add_heading(doc, "5 指标定义", 1)
add_table(doc, ["指标", "定义", "解释"], [
    ("coverage@5", "|Top5 与 truth 的交集| / |truth|", "Recall@5；本报告的课程图 HR@5 口径"),
    ("hit_rate@5", "Top5 至少命中一项时为 1，否则为 0", "标准二值命中率"),
    ("NDCG@5", "二值相关性，理想长度 min(5, |truth|)", "考虑命中位置"),
    ("precision@5", "Top5 命中数 / 5", "每个目标最多 5 个推荐"),
], widths=[1.05, 2.45, 2.27])

add_heading(doc, "6 1,000 条测试集结果", 1)
doc.add_paragraph("四种方法均在同一批 1,000 个合成目标上评测。下表为宏平均，所有方法使用相同的目标、合成真值和 14,201 个候选作品。")
add_table(doc, ["排名", "方法", "目标数", "coverage@5\nRecall", "NDCG@5", "hit_rate@5\n标准", "precision@5"], [
    ("1", "A 五维加权", "1,000", "30.26%", "0.3288", "67.50%", "30.26%"),
    ("2", "D 标签与语义融合", "1,000", "25.80%", "0.2855", "63.20%", "25.80%"),
    ("3", "B 逻辑回归", "1,000", "21.92%", "0.2371", "54.90%", "21.92%"),
    ("4", "C BGE 语义", "1,000", "13.54%", "0.1529", "42.50%", "13.54%"),
], widths=[0.45, 1.55, 0.62, 0.76, 0.76, 0.76, 0.87])
doc.add_paragraph("在该合成协议下，A 的四项总体指标均高于 D，因此本 benchmark 不支持自动用 D 替换 A。D 仍可作为带有语义补充的候选路线，但不能把本次合成结果直接表述为 MAL 或真实用户效果。")

add_heading(doc, "7 系统技术路线和复现", 1)
doc.add_paragraph("运行命令：python generate_synthetic_1000.py --output results/replay_synthetic_1000；随后执行 python run_checks.py --output results/replay_synthetic_1000 和 python verify_delivery.py。重放写入独立目录，不覆盖发布的 results/synthetic_1000。发布目录包含 testset.jsonl、ground_truth_used.jsonl、predictions.jsonl、per_query_metrics.jsonl、metrics.csv、summary.json 和 testset_manifest.json。")
doc.add_paragraph("引擎加载目录、逻辑回归模型和 BGE 向量，针对每个目标一次计算四种分数并进行稳定 Top-5 排序。每个推荐项保存 anime_id、标题、分数和五维特征，逐目标指标可由 ranking_metrics 复算。")
doc.add_paragraph("验证命令：python run_checks.py --output results/replay_synthetic_1000。当前测试覆盖 1,000 条标签、4,000 条方法预测、指标复算、目标自排除、manifest 哈希和离线推荐重放。")

add_heading(doc, "8 结果可信度和限制", 1)
add_bullet(doc, "合成 oracle 与输入字段存在结构关系，A/D 可能受益；参数经过首轮试跑后校准，指标只适用于本协议，不能外推为真实用户效果。")
add_bullet(doc, "MAL 推荐和用户行为没有参与本版本标签，无法回答真实用户是否会点击或喜欢推荐。")
add_bullet(doc, "五维字段无法表达剧情、画风、叙事、节奏和个人偏好；缺失字段也会降低区分度。")
add_bullet(doc, "部署前应建立独立人工标注或真实用户离线集，预先固定抽样、真值和替换门槛，并报告置信区间。")
add_bullet(doc, "合成目标与既有弱监督训练端点可能重叠，当前结果用于回归测试，不承担独立泛化验证。")
add_bullet(doc, "本报告只引用当前 1,000 条合成测试集的结果；独立人工标注或真实用户离线评测仍需另行建立。")

footer = section.footer
footer_paragraph = footer.paragraphs[0]
footer_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
add_page_number(footer_paragraph)

doc.save(OUT)
print(OUT)
